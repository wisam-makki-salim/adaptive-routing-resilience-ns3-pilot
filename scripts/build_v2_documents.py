#!/usr/bin/env python3
"""Build the supervisor brief and technical report from verified v2 results."""

from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)
NAVY = "193A5A"
PALE = "EAF1F7"
GRAY = "D9D9D9"


def set_metadata(doc, title, subject):
    props = doc.core_properties
    props.author = "Wisam Makki Salim"
    props.last_modified_by = "Wisam Makki Salim"
    props.title = title
    props.subject = subject
    props.comments = ""
    props.keywords = "adaptive routing; network resilience; OLSR; ns-3"
    timestamp = datetime(2026, 10, 22, tzinfo=UTC)
    props.created = timestamp
    props.modified = timestamp


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), GRAY)
        borders.append(element)
    tbl_pr.append(borders)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.60)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].paragraph_format.space_after = Pt(4.5)
    styles["Normal"].paragraph_format.line_spacing = 1.05
    for name, size, before, after in (
        ("Title", 19, 0, 6),
        ("Heading 1", 13, 8, 3),
        ("Heading 2", 10.5, 5, 2),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Wisam Makki Salim  |  Pilot Study v2  |  October 2026")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 90, 90)


def add_title(doc, title, subtitle=None):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitle)
        r.bold = True
        r.font.size = Pt(10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Wisam Makki Salim  |  MSc Computer Engineering  |  ORCID 0009-0000-6998-3912")
    r.font.size = Pt(8.5)


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_table(doc, headers, rows, widths=None, font_size=8.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(font_size)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            if row_index % 2:
                shade(cell, PALE)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(text))
            r.font.size = Pt(font_size)
            if widths:
                cell.width = Inches(widths[i])
    return table


def add_figure(doc, relative_path, caption, width=6.65):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ROOT / relative_path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(4)
    r = c.add_run(caption)
    r.italic = True
    r.font.size = Pt(8)


def build_brief():
    doc = Document()
    set_metadata(
        doc,
        "Supervisor Ready PhD Research Brief",
        "Adaptive OLSR resilience and control cost pilot study",
    )
    configure(doc)
    add_title(
        doc,
        "Supervisor Ready PhD Research Brief",
        "When Should Adaptive Network Recovery Be Triggered",
    )
    add_heading(doc, "Research context")
    add_body(
        doc,
        "Fast routing control can shorten disruption, but periodic HELLO and topology messages consume wireless capacity and may destabilize service. Prior work already establishes adaptive MANET routing and adaptive HELLO messaging. The practical question here is narrower: when should faster control be active, and at what cost?",
    )
    add_heading(doc, "Research question")
    add_body(
        doc,
        "Under stable, progressive, abrupt, and repeated disturbances, when does state-triggered OLSR timer adaptation improve packet delivery relative to static OLSR and always-on fast control, and what control and harmful-tail costs accompany any improvement?",
    )
    add_heading(doc, "Existing implementation")
    add_body(
        doc,
        "The ns-3.47 implementation uses a 16-node wireless ad hoc topology and two UDP flows. B0 uses standard OLSR timers. B1 keeps HELLO at 0.5 s and TC at 1.0 s. B2 applies the same fast timers after a signal or transmission-failure trigger. The v2 study added a repeated-disturbance condition and fixed the disturbed node across mechanisms for each seed.",
    )
    add_table(
        doc,
        ["Condition", "Meaning", "Runs per mechanism"],
        [
            ("A", "Stable negative control", "30"),
            ("B", "Progressive mobility disturbance", "30"),
            ("C", "Abrupt next-hop radio failure", "30"),
            ("D", "Repeated out-and-back mobility", "30"),
        ],
        [0.75, 4.75, 1.2],
    )
    add_heading(doc, "Evidence produced")
    add_body(
        doc,
        "The package contains 360 fresh paired runs, per-run logs, raw and processed tables, deterministic bootstrap intervals, three regenerated figures, and explicit checks for missing rows and disturbance-target mismatches.",
    )
    doc.add_page_break()
    add_heading(doc, "Verified findings")
    add_table(
        doc,
        ["Condition", "Always-on PDR effect", "Triggered PDR effect", "Always-on cost", "Triggered cost"],
        [
            ("Stable", "+0.571 [-0.629, +2.304]", "+0.432 [0.000, +1.296]", "+0.05410", "+0.00004"),
            ("Progressive", "+1.591 [+1.214, +1.969]", "+0.108 [-0.069, +0.320]", "+0.05305", "+0.00307"),
            ("Failure", "+1.453 [+1.054, +1.796]", "-0.117 [-0.279, +0.010]", "+0.05043", "+0.00111"),
            ("Repeated", "+3.392 [+2.738, +4.044]", "+0.404 [-0.260, +1.050]", "+0.05357", "+0.00378"),
        ],
        [1.05, 1.75, 1.75, 1.1, 1.1],
        7.6,
    )
    add_body(
        doc,
        " PDR effects are paired mean percentage-point differences from static OLSR with 95% percentile-bootstrap intervals. Cost is normalized OLSR control bytes per delivered payload byte.",
        "Reading the table.",
    )
    add_figure(
        doc,
        "figures/v2/v2_benefit_cost_tradeoff.png",
        "Figure 1  Mean delivery effect and normalized control cost relative to static OLSR. Error bars show 95% paired bootstrap intervals.",
        5.35,
    )
    add_heading(doc, "Interpretation and limitations")
    add_body(
        doc,
        "Always-on fast control improves mean delivery in all three disturbance conditions, but adds a large and consistent control cost and raises p95 delay in most runs. The locked triggered rule does not reproduce that benefit reliably. It also triggers in 3 of 30 stable runs, so the 20-event threshold is not validated. One topology family, two flows, deterministic path loss, and an aggregate recovery definition limit generalization.",
    )
    add_heading(doc, "Proposed doctoral extension")
    add_body(
        doc,
        "The next step is a guarded trigger and exit policy that combines route state with channel-load evidence, caps reactive duration, and treats worst-flow harm as a constraint. Confirmation requires fresh development and evaluation seeds, several topology and mobility families, heterogeneous traffic, and selected testbed checks.",
    )
    add_heading(doc, "Repository and feasibility")
    add_body(
        doc,
        "Repository: https://github.com/wisam-makki-salim/adaptive-routing-resilience-ns3-pilot. The work is suitable for an externally conducted doctorate because simulation, analysis, and documentation can be performed remotely; short visits could support design review and physical validation. Any formal arrangement remains subject to the supervisor and university regulations.",
    )
    path = OUT / "Wisam_Makki_Salim_Supervisor_Ready_Research_Brief_v2.docx"
    doc.save(path)
    return path


def build_report():
    doc = Document()
    set_metadata(
        doc,
        "Pilot Study v2 Technical Report",
        "Resilience and control cost tradeoffs in adaptive OLSR",
    )
    configure(doc)
    add_title(
        doc,
        "Pilot Study v2 Technical Report",
        "Resilience and Control Cost Tradeoffs in Adaptive OLSR",
    )
    add_heading(doc, "Abstract")
    add_body(
        doc,
        "This study asks whether state-triggered OLSR timer adaptation provides a better resilience-cost balance than static operation and permanent fast control. A fresh ns-3.47 experiment executed 360 paired runs across stable, progressive, abrupt-failure, and repeated-disturbance conditions. Always-on fast control improved mean packet delivery in the three disturbance conditions but increased normalized control overhead by about 0.050 to 0.054 and usually increased p95 delay. The pre-specified state-triggered rule did not produce a repeatable delivery gain and activated in 3 of 30 stable runs. The evidence therefore rejects the tested trigger as a validated solution while identifying a specific doctoral problem: guarded activation and exit control under explicit overhead and worst-flow constraints.",
    )
    add_heading(doc, "Research question and contribution boundary")
    add_body(
        doc,
        "The primary question is when state-triggered fast control improves delivery relative to static and always-on operation, and what costs accompany any gain. The contribution is a reproducible three-way comparison with a fixed disturbance target, fresh paired seeds, and retained negative results. Timer adaptation is established prior work; this report does not claim a new routing protocol or state-of-the-art performance.",
    )
    add_heading(doc, "Related work position")
    add_body(
        doc,
        "RFC 3626 defines periodic OLSR HELLO and TC messaging. Qin and Kunz studied adaptive MANET routing, Hoebeke and colleagues examined context-dependent routing modes, and Han and Lee studied adaptive HELLO messaging. These studies establish the broader direction. The pilot addresses a narrower design question through an explicit static, always-on, and state-triggered comparison.",
    )
    add_figure(
        doc,
        "figures/v2/v2_experimental_architecture.png",
        "Figure 1  Executed v2 workflow. Every condition uses the same 30 run numbers across B0, B1, and B2.",
        6.7,
    )
    add_heading(doc, "Methods")
    add_body(
        doc,
        "Sixteen nodes form a jittered four-by-four IEEE 802.11g ad hoc topology. Two UDP flows transmit 512-byte packets every 20 ms during a 120 s simulation. B0 uses HELLO 2.0 s and TC 5.0 s. B1 uses HELLO 0.5 s and TC 1.0 s throughout. B2 enters the same fast mode at -81 dBm after three weak samples or after 20 final transmission failures within the implemented burst rule.",
    )
    add_table(
        doc,
        ["Item", "Locked value"],
        [
            ("Simulator", "ns-3.47, optimized build"),
            ("Fresh runs", "201–230, paired across mechanisms"),
            ("Conditions", "A stable; B progressive; C failure; D repeated"),
            ("Primary outcome", "Packet delivery ratio"),
            ("Guardrails", "Worst-flow PDR, p95 delay, overhead, recovery, route availability"),
            ("Uncertainty", "50,000 deterministic paired bootstrap resamples"),
        ],
        [1.65, 5.05],
    )
    add_body(
        doc,
        "For each fresh run, a static B0 probe identifies the active next-hop target. That node is then fixed for B0, B1, and B2. This correction prevents the mechanism from changing the disturbance it receives. Every simulator process writes a separate log and a 22-field result row. The analysis rejects incomplete matrices or target mismatches.",
    )
    add_heading(doc, "Results")
    add_table(
        doc,
        ["Condition", "B1 PDR difference", "B2 PDR difference", "B1 overhead", "B2 overhead"],
        [
            ("A", "+0.571 [-0.629, +2.304]", "+0.432 [0.000, +1.296]", "+0.05410", "+0.00004"),
            ("B", "+1.591 [+1.214, +1.969]", "+0.108 [-0.069, +0.320]", "+0.05305", "+0.00307"),
            ("C", "+1.453 [+1.054, +1.796]", "-0.117 [-0.279, +0.010]", "+0.05043", "+0.00111"),
            ("D", "+3.392 [+2.738, +4.044]", "+0.404 [-0.260, +1.050]", "+0.05357", "+0.00378"),
        ],
        [0.7, 1.8, 1.8, 1.15, 1.15],
        7.7,
    )
    add_figure(
        doc,
        "figures/v2/v2_resilience_performance.png",
        "Figure 2  Mean paired PDR difference from static OLSR with 95% bootstrap intervals.",
        6.45,
    )
    add_body(
        doc,
        "B1 increased PDR in 28/30 B runs, 28/30 C runs, and 29/30 D runs. B2 effects were mixed: 14/30 B runs improved, 8/30 C runs improved, and 17/30 D runs improved. Relative to B1, B2 PDR was lower by 1.483 points in B, 1.570 in C, and 2.988 in D. The corresponding intervals were wholly below zero.",
    )
    add_figure(
        doc,
        "figures/v2/v2_benefit_cost_tradeoff.png",
        "Figure 3  Mean resilience effect and normalized control cost. Labels A–D denote the four conditions.",
        5.8,
    )
    add_heading(doc, "Stable control and trigger behavior")
    add_body(
        doc,
        "B2 produced four failure-trigger transitions across three stable runs and no recovery transition. Twenty-nine of 30 stable PDR pairs matched B0 exactly; one outlier produced the positive mean. This violates H1 and shows that the pre-specified 20-event threshold is too sensitive for the intended negative-control behavior. B2 triggered in all B and D runs and in 27/30 C runs.",
    )
    add_heading(doc, "Worst flow and recovery")
    add_body(
        doc,
        "B1 increased mean minimum-flow PDR by 3.697 points in B, 3.349 in C, and 7.229 in D. B2 changes were +0.217, -0.220, and +0.475, with intervals crossing zero. The recovery field uses -2 for non-recovery and was not averaged as a duration. All B and C runs met the operational definition. In D, B1 did not recover in 6/30 runs; B0 and B2 did so in all 30. The zero-second median in D also shows that the aggregate recovery definition is too coarse for repeated events.",
    )
    add_heading(doc, "Threats to validity")
    add_bullets(
        doc,
        [
            "One topology family and deterministic log-distance path loss limit external validity.",
            "Two constant-rate flows do not represent heterogeneous IoT traffic.",
            "The disturbed relay is anchored fairly but still comes from one source route.",
            "The controller's recovery depends on renewed signal evidence and may remain reactive after route change.",
            "Bootstrap intervals describe variation across these 30 runs; they do not prove generality.",
            "No energy claim is made because no energy model is part of the v2 experiment.",
        ],
    )
    add_heading(doc, "Independent review")
    add_body(
        doc,
        "The comparisons are suitable for supervisor review after the target-selection confound and recovery sentinel were corrected. The tested B2 controller is not suitable for a superiority claim. The project is strongest as proof of execution: it exposes a failed trigger, shows the attainable benefit and cost of fast control, and leaves a precise research problem for doctoral work.",
    )
    add_heading(doc, "Doctoral research path")
    add_bullets(
        doc,
        [
            "Design a guarded trigger and exit policy with a finite reactive budget, route-state confirmation, and channel-load safeguards.",
            "Validate across topology, density, mobility, fading, traffic, and failure-location blocks using fresh seeds.",
            "Select policies under explicit worst-flow, tail-loss, delay, and control-cost constraints, then test selected cases on a small physical platform.",
        ],
    )
    add_heading(doc, "Reproducibility statement")
    add_body(
        doc,
        "The repository retains the locked YAML protocol, C++ source, execution and analysis scripts, raw 360-run table, target anchors, per-run logs, processed comparisons, recovery and trigger summaries, and figure sources. The exact command is documented in the README. The main result table and PNG figures regenerate deterministically from the stored raw CSV.",
    )
    add_heading(doc, "References")
    refs = [
        "Clausen T, Jacquet P. Optimized Link State Routing Protocol. RFC 3626. 2003. doi:10.17487/RFC3626.",
        "Qin L, Kunz T. Adaptive MANET Routing: A Case Study. ADHOC-NOW 2008:43–57. doi:10.1007/978-3-540-85209-4_4.",
        "Hoebeke J et al. Adaptive routing for mobile ad hoc networks. EURASIP Journal on Wireless Communications and Networking. 2012;2012:126. doi:10.1186/1687-1499-2012-126.",
        "Han SY, Lee D. An Adaptive Hello Messaging Scheme for Neighbor Discovery in On-Demand MANET Routing Protocols. IEEE Communications Letters. 2013;17(5):1040–1043. doi:10.1109/LCOMM.2013.040213.130076.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.add_run(ref)
    path = OUT / "Wisam_Makki_Salim_Pilot_Study_v2_Technical_Report.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_brief())
    print(build_report())
