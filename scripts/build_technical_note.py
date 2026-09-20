#!/usr/bin/env python3
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Wisam_Makki_Salim_Adaptive_Routing_Pilot_Technical_Note.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def cell_margins(cell, top=55, start=70, bottom=55, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, header_fill="17365D", font_size=7.4):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_border(cell)
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0:
                set_cell_shading(cell, header_fill)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "EEF3F8")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
                    if r_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = False
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(7.8)
    return p


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.52)
section.bottom_margin = Inches(0.48)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)
section.header_distance = Inches(0.22)
section.footer_distance = Inches(0.22)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(9.3)
styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
styles["Normal"].paragraph_format.space_after = Pt(3.2)
styles["Normal"].paragraph_format.line_spacing = 1.02
for name, size, before, after in (("Title", 19, 0, 5), ("Subtitle", 12, 0, 8),
                                  ("Heading 1", 12, 6, 2), ("Heading 2", 10, 4, 1)):
    st = styles[name]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.bold = name != "Subtitle"
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

# Remove Word/LibreOffice theme borders from the title style.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

doc.settings.odd_and_even_pages_header_footer = True

for header_part in (section.header, section.even_page_header):
    header = header_part.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("REPRODUCIBLE NETWORKING PILOT  |  SEPTEMBER 2026")
    run.font.name = "Arial"; run.font.size = Pt(7); run.font.color.rgb = RGBColor(90, 90, 90)
for footer_part in (section.footer, section.even_page_footer):
    footer = footer_part.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Wisam Makki Salim  |  Adaptive Routing Pilot  |  ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for r in footer.runs:
        r.font.name = "Arial"; r.font.size = Pt(7); r.font.color.rgb = RGBColor(90, 90, 90)

# Page 1
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Adaptive Routing under Mobility and Failures")
subtitle = doc.add_paragraph(style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("A Reproducible ns-3 Pilot with Robustness and Component Ablation")
author = doc.add_paragraph()
author.alignment = WD_ALIGN_PARAGRAPH.CENTER
author.paragraph_format.space_after = Pt(7)
r = author.add_run("Wisam Makki Salim\nComputer Engineering Researcher  |  Al-Iraqia University, Baghdad, Iraq\nORCID 0009-0000-6998-3912")
r.bold = True; r.font.size = Pt(8.4)

heading(doc, "Abstract")
body(doc, "This study tests whether lightweight OLSR timer adaptation can preserve service under mobility, link degradation, and node failure. An ns-3.47 experiment compares standard OLSR with a controller that reacts to weak signals and final transmission failures. The evaluation uses a 16-node wireless ad hoc topology, two UDP flows, three controlled scenarios, and 20 paired runs per comparison. The results do not support the original superiority hypothesis: service effects vary across runs, while routing overhead rises whenever adaptation triggers. Component ablation identifies accelerated HELLO messaging as the source of an adverse cross-flow tail in the tested configuration. A lower failure threshold shows exploratory gains but still requires independent evaluation.")

heading(doc, "Research Question and Hypothesis")
body(doc, "Can a lightweight adaptive routing-control mechanism improve service resilience under changing mobility, link degradation, and failure conditions compared with a static baseline, without introducing excessive latency or routing overhead?")
body(doc, "The working hypothesis predicted shorter disruption and recovery with a modest control cost. The experiment tests this claim rather than assuming that faster protocol control is beneficial.")

heading(doc, "Experimental Design")
table = doc.add_table(rows=1, cols=2)
table.columns[0].width = Inches(1.65); table.columns[1].width = Inches(5.35)
for i, text in enumerate(("Element", "Locked specification")): table.rows[0].cells[i].text = text
design_rows = [
    ("Simulator", "ns-3.47; IEEE 802.11g ad hoc; standard OLSR baseline"),
    ("Network", "16-node jittered 4 by 4 grid; two concurrent UDP flows; 512-byte packets every 20 ms"),
    ("Adaptation", "EWMA link signal and final transmission failures; reactive HELLO 0.5 s and TC 1.0 s"),
    ("Scenarios", "A normal operation; B gradual next-hop displacement; C active next-hop radio failure"),
    ("Evaluation", "Runs 101 to 120 paired by seed; 50,000 deterministic bootstrap resamples"),
    ("Measures", "PDR, worst-flow PDR, delay, goodput, control overhead, recovery, and route stability"),
]
for left, right in design_rows:
    cells = table.add_row().cells; cells[0].text = left; cells[1].text = right
style_table(table, font_size=7.6)

body(doc, "The controller enters a reactive state after three weak-signal samples at or below -81 dBm or 40 final transmission failures within one second. Hysteresis and a ten-second hold-down govern reversion. Runs are retained regardless of direction or magnitude.")

# Page 2
page_break(doc)
heading(doc, "Locked Evaluation Results")
body(doc, "Scenario A acts as a negative control: baseline and adaptive outputs are identical in all 20 pairs and no trigger fires. In Scenario B, the PDR interval crosses zero while normalized overhead increases in every pair. In Scenario C, the median PDR effect is zero, but run 102 reveals a severe adverse tail that is reproduced exactly with diagnostic tracing.")

table = doc.add_table(rows=1, cols=5)
widths = [0.55, 1.25, 1.5, 1.5, 2.2]
for i, w in enumerate(widths): table.columns[i].width = Inches(w)
headers = ["Case", "Median PDR baseline", "Median PDR adaptive", "Mean PDR difference (95% CI)", "Evidence" ]
for i, text in enumerate(headers): table.rows[0].cells[i].text = text
rows = [
    ("A", "100.00%", "100.00%", "0.000  [0.000, 0.000]", "No trigger; identical outputs"),
    ("B", "97.10%", "97.06%", "+0.057  [-0.119, +0.312]", "No stable benefit; overhead higher in 20 of 20"),
    ("C", "97.07%", "97.21%", "-0.659  [-2.777, +0.604]", "Typical effect near zero; adverse tail retained"),
]
for row in rows:
    cells = table.add_row().cells
    for i, text in enumerate(row): cells[i].text = text
style_table(table, font_size=7.1)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(1)
p.add_run().add_picture(str(ROOT / "figures/g5_paired_pdr_differences.png"), width=Inches(6.65))
caption(doc, "Figure 1  Adaptive minus baseline PDR for every paired final evaluation run")

heading(doc, "Interpretation")
body(doc, "The Scenario B mean effect is +0.057 percentage points with a 95 percent bootstrap interval from -0.119 to +0.312. Adaptive PDR is lower in 10 pairs, equal in 6, and higher in 4. The mean overhead difference is +0.002518 with an interval wholly above zero. The mechanism therefore incurs a repeatable cost without a repeatable service benefit under gradual degradation.")
body(doc, "In Scenario C, 9 pairs improve, 4 are equal, and 7 decline. Run 102 reduces PDR from 95.51 percent to 76.34 percent and extends recovery from 7 s to 36 s. The rerun matches the stored result exactly. Per-flow tracing shows that one flow recovers while another remains largely undelivered despite a source route.")

# Page 3
page_break(doc)
heading(doc, "Component Ablation")
body(doc, "Scenario C was rerun with detection-only, HELLO-only, and topology-control-only variants. Detection-only matches baseline in every network field across all 20 runs. HELLO-only matches the full adaptive mechanism exactly. Topology-control-only is nearly identical to baseline. The evidence therefore attributes the measured cost and adverse tail to shortening HELLO from 2.0 s to 0.5 s.")

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
p.add_run().add_picture(str(ROOT / "figures/g6_component_ablation.png"), width=Inches(6.65))
caption(doc, "Figure 2  Scenario C component ablation with paired mean effects and bootstrap intervals")

table = doc.add_table(rows=1, cols=5)
for i, w in enumerate([1.35, 1.05, 1.05, 1.15, 2.35]): table.columns[i].width = Inches(w)
for i, text in enumerate(["Run 102 variant", "PDR", "Overhead", "Recovery", "Result"]): table.rows[0].cells[i].text = text
for row in [
    ("Baseline", "95.51%", "0.017424", "7 s", "Reference"),
    ("Detection only", "95.51%", "0.017424", "7 s", "Exact baseline match"),
    ("HELLO only", "76.34%", "0.023060", "36 s", "Exact full-adaptive match"),
    ("TC only", "95.51%", "0.017424", "7 s", "Exact baseline match"),
    ("Full adaptive", "76.34%", "0.023060", "36 s", "Adverse tail"),
]:
    cells = table.add_row().cells
    for i, text in enumerate(row): cells[i].text = text
style_table(table, font_size=7.2)

heading(doc, "Causal Boundary")
body(doc, "Within this simulation, the ablation links the run 102 outcome to accelerated HELLO messaging. It does not establish the lower-level mechanism. The available traces cannot distinguish transient routing-state inconsistency from wireless contention or another OLSR interaction. That distinction requires additional instrumentation rather than inference from aggregate delivery.")

# Page 4
page_break(doc)
heading(doc, "Threshold Sensitivity")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
p.add_run().add_picture(str(ROOT / "figures/g6_threshold_sensitivity.png"), width=Inches(6.35))
caption(doc, "Figure 3  Threshold sensitivity with paired mean PDR effects and bootstrap intervals")

body(doc, "A failure threshold of 20 events produces a mean PDR gain of +0.434 percentage points [95 percent CI +0.082, +0.884] and a mean recovery reduction of 0.45 s. Its leave-one-out PDR mean remains positive, but overhead increases in every run. This is an exploratory redesign candidate, not a confirmed result, because it reuses the evaluation seeds and topology. A threshold of 80 never triggers and only reproduces baseline behavior.")
body(doc, "In Scenario B, -78 dBm produces 317 signal transitions and a large overhead increase. At -84 dBm, signal triggering disappears and the failure channel triggers in 19 runs. The channels substitute for one another, so these data do not define an optimal RSSI threshold.")

heading(doc, "Limitations and Doctoral Direction")
body(doc, "The pilot uses one topology family, one density, two constant-rate flows, and an idealized failure. It does not measure channel occupancy or route loops directly. Aggregate recovery can conceal flow-level unfairness, and bootstrap intervals over 20 paired runs do not establish generality across network environments.")
body(doc, "A doctoral extension should replace unconditional timer acceleration with guarded control that combines link risk, route state, and congestion evidence; limits the duration and scope of HELLO changes; and treats cross-flow harm as a first-class outcome. Confirmation requires fresh seeds, multiple mobility and topology families, heterogeneous traffic, and explicit tail-risk and fairness measures. Adversarial disruption can then be introduced as a separate controlled factor.")

heading(doc, "Conclusion")
body(doc, "The original timer-switching mechanism does not provide reliable evidence of improved resilience. It raises overhead and can create severe cross-flow harm after failure. The resulting code, controlled comparisons, uncertainty analysis, diagnostic traces, and component ablation provide a reproducible basis for studying safer adaptive control.")

heading(doc, "References")
for ref in [
    "1. ns-3 Consortium. ns-3 Network Simulator Documentation. https://www.nsnam.org/documentation/",
    "2. Clausen T and Jacquet P. Optimized Link State Routing Protocol OLSR. RFC 3626. IETF 2003. https://www.rfc-editor.org/rfc/rfc3626",
    "3. German Research Foundation. Guidelines for Safeguarding Good Research Practice. https://www.dfg.de/en/research-funding/funding-principles/good-scientific-practice",
]:
    p = doc.add_paragraph(ref)
    p.paragraph_format.space_after = Pt(1)
    for run in p.runs: run.font.size = Pt(7.2)

doc.core_properties.title = "Adaptive Routing under Mobility and Failures"
doc.core_properties.subject = "Reproducible ns-3 study of adaptive OLSR control"
doc.core_properties.author = "Wisam Makki Salim"
doc.core_properties.keywords = "ns-3, OLSR, adaptive routing, resilient networking, reproducibility"
doc.core_properties.comments = ""
doc.save(OUT)
print(OUT)
